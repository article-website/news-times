"""
Pembuat slide presentasi NewsTimes.

Berkas ini adalah SUMBER KEBENARAN untuk isi presentasi. Jangan mengedit
file .pptx langsung kalau ingin perubahannya bertahan — ubah di sini, lalu
jalankan ulang:

    pip install python-pptx
    python docs/presentasi/buat-slide.py

Hasilnya: docs/presentasi/NewsTimes.pptx

Catatan bicara tiap slide ikut tertanam sebagai "speaker notes" di PowerPoint.
Angka dan status di sini harus cocok dengan STATUS.md. Kalau STATUS.md berubah,
perbarui juga berkas ini.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

# --- Warna, mengikuti tampilan situs -----------------------------------------
GELAP = RGBColor(0x0F, 0x17, 0x2A)   # slate-900, sama dengan footer situs
BIRU = RGBColor(0x25, 0x63, 0xEB)    # blue-600, warna aksen situs
PUTIH = RGBColor(0xFF, 0xFF, 0xFF)
TEKS = RGBColor(0x1F, 0x29, 0x37)
REDUP = RGBColor(0x6B, 0x72, 0x80)
GARIS = RGBColor(0xE5, 0xE7, 0xEB)
HIJAU = RGBColor(0x05, 0x96, 0x69)
MERAH = RGBColor(0xDC, 0x26, 0x26)
KUNING = RGBColor(0xB4, 0x53, 0x09)

LEBAR = Inches(13.333)
TINGGI = Inches(7.5)
TEPI = Inches(0.9)
ISI = LEBAR - 2 * TEPI

prs = Presentation()
prs.slide_width = LEBAR
prs.slide_height = TINGGI
KOSONG = prs.slide_layouts[6]


def slide_baru(latar=PUTIH):
    s = prs.slides.add_slide(KOSONG)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = latar
    return s


def teks(slide, kiri, atas, lebar, tinggi, isi, ukuran=18, warna=TEKS,
         tebal=False, rata=PP_ALIGN.LEFT, spasi=Pt(6)):
    kotak = slide.shapes.add_textbox(kiri, atas, lebar, tinggi)
    tf = kotak.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    baris = isi if isinstance(isi, list) else [isi]
    for i, b in enumerate(baris):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = b
        p.alignment = rata
        p.space_after = spasi
        p.font.size = Pt(ukuran)
        p.font.bold = tebal
        p.font.color.rgb = warna
        p.font.name = "Segoe UI"
    return kotak


def garis_aksen(slide, atas=Inches(1.55), lebar=Inches(1.1), warna=BIRU):
    bar = slide.shapes.add_shape(1, TEPI, atas, lebar, Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = warna
    bar.line.fill.background()
    bar.shadow.inherit = False


def judul_halaman(slide, judul, kicker=None):
    if kicker:
        teks(slide, TEPI, Inches(0.55), ISI, Inches(0.3), kicker.upper(),
             ukuran=11, warna=BIRU, tebal=True)
    teks(slide, TEPI, Inches(0.9), ISI, Inches(0.6), judul, ukuran=30, tebal=True)
    garis_aksen(slide)


def catatan(slide, isi):
    slide.notes_slide.notes_text_frame.text = isi


def tabel(slide, kepala, baris, atas, lebar_kolom, ukuran=13):
    n_baris = len(baris) + 1
    n_kolom = len(kepala)
    total = sum(lebar_kolom)
    bentuk = slide.shapes.add_table(n_baris, n_kolom, TEPI, atas, total,
                                    Inches(0.4) * n_baris)
    t = bentuk.table
    t.first_row = True
    for i, w in enumerate(lebar_kolom):
        t.columns[i].width = Emu(int(w))

    for c, judul_kolom in enumerate(kepala):
        sel = t.cell(0, c)
        sel.text = judul_kolom
        sel.fill.solid()
        sel.fill.fore_color.rgb = GELAP
        p = sel.text_frame.paragraphs[0]
        p.font.size = Pt(ukuran)
        p.font.bold = True
        p.font.color.rgb = PUTIH
        p.font.name = "Segoe UI"

    for r, isi_baris in enumerate(baris, start=1):
        for c, isi_sel in enumerate(isi_baris):
            nilai, warna = isi_sel if isinstance(isi_sel, tuple) else (isi_sel, TEKS)
            sel = t.cell(r, c)
            sel.text = nilai
            sel.fill.solid()
            sel.fill.fore_color.rgb = PUTIH if r % 2 else RGBColor(0xF9, 0xFA, 0xFB)
            p = sel.text_frame.paragraphs[0]
            p.font.size = Pt(ukuran)
            p.font.color.rgb = warna
            p.font.bold = warna in (HIJAU, MERAH, KUNING)
            p.font.name = "Segoe UI"
    return t


def poin(slide, daftar, atas=Inches(2.0), ukuran=18, jarak=Inches(0.72)):
    for i, (utama, penjelas) in enumerate(daftar):
        y = atas + jarak * i
        titik = slide.shapes.add_shape(9, TEPI, y + Inches(0.09), Pt(9), Pt(9))
        titik.fill.solid()
        titik.fill.fore_color.rgb = BIRU
        titik.line.fill.background()
        titik.shadow.inherit = False
        teks(slide, TEPI + Inches(0.32), y, ISI - Inches(0.32), Inches(0.3),
             utama, ukuran=ukuran, tebal=True)
        if penjelas:
            teks(slide, TEPI + Inches(0.32), y + Inches(0.31),
                 ISI - Inches(0.32), Inches(0.3), penjelas,
                 ukuran=ukuran - 4, warna=REDUP)


# =============================================================================
# 1 — Sampul
# =============================================================================
s = slide_baru(GELAP)
bar = s.shapes.add_shape(1, TEPI, Inches(2.35), Inches(1.2), Pt(5))
bar.fill.solid()
bar.fill.fore_color.rgb = BIRU
bar.line.fill.background()
bar.shadow.inherit = False
teks(s, TEPI, Inches(1.55), ISI, Inches(0.4), "LAPORAN PROGRES",
     ukuran=13, warna=BIRU, tebal=True)
teks(s, TEPI, Inches(2.6), ISI, Inches(1.0), "NewsTimes", ukuran=54,
     warna=PUTIH, tebal=True)
teks(s, TEPI, Inches(3.7), Inches(8.5), Inches(0.8),
     "Portal berita dengan ruang redaksi sendiri", ukuran=22,
     warna=RGBColor(0xCB, 0xD5, 0xE1))
teks(s, TEPI, Inches(5.6), ISI, Inches(0.9),
     ["Tim 5 orang  ·  Next.js 16  ·  PostgreSQL di Neon",
      "11 September 2026"],
     ukuran=14, warna=REDUP, spasi=Pt(4))
catatan(s, "Perkenalan singkat. Sebut ini laporan progres, bukan demo produk "
           "jadi — supaya harapan pendengar pas sejak awal.")

# =============================================================================
# 2 — Ringkasan
# =============================================================================
s = slide_baru()
judul_halaman(s, "Kalau cuma sempat dengar satu slide", "ringkasan")
poin(s, [
    ("Fondasi datanya sudah selesai dan terbukti jalan",
     "5 tabel, 4 kumpulan fungsi akses data, 163 pengecekan otomatis lulus"),
    ("Tapi belum ada satu halaman pun yang memakainya",
     "Semua halaman masih membaca 5 artikel yang ditulis tangan di dalam kode"),
    ("Redaksi belum bisa menerbitkan artikel sendiri",
     "Halaman admin sudah ada, tapi datanya masih tersimpan di browser"),
    ("Yang paling mendesak: dua peran tim masih kosong",
     "Tanpa Orang 5, branch utama tidak terkunci dan tidak ada pengecekan otomatis"),
], atas=Inches(2.05), jarak=Inches(1.05))
catatan(s, "Ini slide paling penting. Kalau presentasi dipotong, empat kalimat "
           "ini yang harus tersampaikan. Jangan menutupi poin kedua — justru "
           "itu yang menunjukkan kita tahu posisi sendiri.")

# =============================================================================
# 3 — Masalah
# =============================================================================
s = slide_baru()
judul_halaman(s, "Kondisi awal: tampilannya jadi, isinya mati", "masalah")
teks(s, TEPI, Inches(1.85), ISI, Inches(0.4),
     "Lima artikel ditulis tangan langsung di dalam kode, di src/data/articles.ts",
     ukuran=15, warna=REDUP)
tabel(s, ["Masalah", "Dampaknya"], [
    ["Menambah artikel harus mengubah kode lalu deploy ulang",
     "Redaksi tidak bisa menulis sendiri — semua lewat programmer"],
    ['Tanggal disimpan sebagai teks: "28 Juni 2025"',
     "Artikel tidak bisa diurutkan dari yang terbaru"],
    ["Enam menu kategori mengarah ke href=\"#\"",
     "Pengunjung tidak bisa menelusuri per topik"],
    ["Tidak ada pencarian",
     "Artikel lama hilang begitu tergeser artikel baru"],
    ["Form newsletter hanya memunculkan alert()",
     "Tidak ada satu pun email yang tersimpan"],
], atas=Inches(2.45), lebar_kolom=[Inches(6.0), Inches(5.5)])
catatan(s, "Tekankan baris kedua — bug tanggal itu yang paling meyakinkan, "
           "karena mengurutkan berita dari yang terbaru adalah fungsi paling "
           "dasar sebuah situs berita.")

# =============================================================================
# 4 — Bug tanggal
# =============================================================================
s = slide_baru(GELAP)
teks(s, TEPI, Inches(0.8), ISI, Inches(0.4), "TEMUAN", ukuran=12,
     warna=BIRU, tebal=True)
teks(s, TEPI, Inches(1.4), ISI, Inches(1.4),
     "Komputer membaca tanggal itu\nsebagai tulisan, bukan tanggal",
     ukuran=34, warna=PUTIH, tebal=True)
teks(s, TEPI, Inches(3.3), Inches(11.0), Inches(1.4),
     ['"28 Juni"  dianggap LEBIH KECIL dari  "3 Mei"',
      "karena diurutkan huruf per huruf, dan '2' datang sebelum '3'"],
     ukuran=22, warna=RGBColor(0xCB, 0xD5, 0xE1), spasi=Pt(10))
teks(s, TEPI, Inches(5.3), Inches(11.0), Inches(1.2),
     ["Artinya artikel tidak bisa diurutkan dari yang terbaru sama sekali.",
      "Sekarang disimpan sebagai tanggal sungguhan, dan penerjemahnya menolak "
      "tanggal yang tidak ada seperti 31 Februari."],
     ukuran=15, warna=REDUP, spasi=Pt(6))
catatan(s, "Slide ini biasanya memancing pertanyaan. Jawabannya: JavaScript "
           "kalau diberi 31 Februari diam-diam menggesernya jadi 3 Maret, jadi "
           "hasil terjemahan dicek balik dan artikel dengan tanggal ngawur "
           "dilewati.")

# =============================================================================
# 5 — Arsitektur
# =============================================================================
s = slide_baru()
judul_halaman(s, "Satu aplikasi, bukan dua", "arsitektur")
teks(s, TEPI, Inches(2.0), ISI, Inches(0.5),
     "Halaman  →  fungsi akses data  →  database", ukuran=26, tebal=True,
     warna=BIRU)
teks(s, TEPI, Inches(2.75), Inches(11.2), Inches(1.0),
     "Tidak ada backend terpisah. Halaman memanggil fungsi biasa di dalam "
     "server yang sama — tidak lewat jaringan, tidak lewat alamat URL.",
     ukuran=16, warna=TEKS)
tabel(s, ["Keputusan", "Alasannya"], [
    ["Monolith, bukan backend terpisah",
     "Tim 5 orang, repo kecil. Dua repo berarti dua deploy dan dua tempat rusak"],
    ["API hanya untuk yang wajib berupa URL",
     "sitemap.xml dan rss.xml dibaca mesin luar. Form cukup Server Action"],
    ["Halaman tidak boleh menyentuh database langsung",
     "Semua lewat 4 kumpulan fungsi, supaya perubahan database tidak merembet"],
], atas=Inches(4.05), lebar_kolom=[Inches(4.3), Inches(7.2)])
catatan(s, "Kalau ditanya 'kenapa tidak pakai API?', jawabannya: API yang "
           "dihindari monolith itu backend sebagai proyek terpisah, bukan "
           "alamat URL. Yang benar-benar wajib URL cuma sitemap dan RSS.")

# =============================================================================
# 6 — Model data
# =============================================================================
s = slide_baru()
judul_halaman(s, "Lima tabel", "model data")
tabel(s, ["Tabel", "Isinya", "Yang perlu diingat"], [
    ["Article", "Artikel", "publishedAt boleh kosong = draft. Punya 3 index"],
    ["Category", "Kategori", "Punya kolom urutan, supaya navbar tidak hardcode"],
    ["Author", "Penulis artikel", "Terpisah dari akun login"],
    ["User", "Akun redaksi", "Hanya menyimpan hash password"],
    ["NewsletterSubscriber", "Pendaftar newsletter", "Berhenti dicatat, data tidak dihapus"],
], atas=Inches(2.1), lebar_kolom=[Inches(2.8), Inches(2.9), Inches(5.8)])
teks(s, TEPI, Inches(5.1), ISI, Inches(0.8),
     "Satu artikel punya tepat satu kategori dan satu penulis. Kategori dan "
     "penulis tidak bisa dihapus selama masih dipakai artikel.",
     ukuran=15, warna=REDUP)
catatan(s, "Kalau ditanya soal index: index itu seperti daftar isi buku. Ada "
           "tiga, masing-masing mengikuti pola pencarian yang paling sering "
           "dipakai halaman.")

# =============================================================================
# 7 — Keputusan: dua versi
# =============================================================================
s = slide_baru()
judul_halaman(s, "Dua versi, supaya tim tidak saling menunggu", "keputusan kunci")
teks(s, TEPI, Inches(1.95), Inches(11.2), Inches(0.8),
     "Kalau halaman langsung menyambung ke database, tidak ada satu pun anggota "
     "tim yang bisa mulai bekerja sebelum databasenya jadi.",
     ukuran=16, warna=TEKS)
tabel(s, ["", "Versi data contoh", "Versi database"], [
    ["Butuh database?", ("Tidak", HIJAU), "Ya"],
    ["Dipakai kapan", "Selama tim membangun tampilan", "Setelah database siap"],
    ["Bentuk hasilnya", ("Sama persis", HIJAU), ("Sama persis", HIJAU)],
    ["Cara menukar", "Satu baris di berkas setelan", "Satu baris di berkas setelan"],
], atas=Inches(3.0), lebar_kolom=[Inches(2.6), Inches(4.4), Inches(4.5)])
teks(s, TEPI, Inches(5.5), ISI, Inches(0.6),
     "Sudah dibuktikan: 20 pemeriksaan membandingkan hasil kedua versi kolom "
     "per kolom — semuanya sama, nol perbedaan.",
     ukuran=15, warna=REDUP)
catatan(s, "Ini keputusan yang paling berdampak ke tim. Kalimat andalannya: "
           "halaman tidak perlu tahu datanya datang dari mana.")

# =============================================================================
# 8 — Keputusan: draft tidak bocor
# =============================================================================
s = slide_baru()
judul_halaman(s, "Draft tidak boleh bocor", "keputusan kunci")
teks(s, TEPI, Inches(2.0), Inches(11.2), Inches(0.5),
     "Artikel terlihat publik hanya kalau DUA syarat terpenuhi:",
     ukuran=17, warna=TEKS)
poin(s, [
    ("Statusnya sudah terbit", ""),
    ("Tanggal terbitnya sudah lewat", ""),
], atas=Inches(2.75), ukuran=20, jarak=Inches(0.6))
teks(s, TEPI, Inches(4.15), Inches(11.2), Inches(1.2),
     ["Kenapa syarat kedua penting?",
      "Tanpa itu, artikel yang dijadwalkan tayang besok sudah bisa dibaca hari "
      "ini oleh siapa pun yang menebak alamatnya. Bocor sebelum waktunya."],
     ukuran=16, warna=TEKS, spasi=Pt(8))
teks(s, TEPI, Inches(5.75), Inches(11.2), Inches(0.6),
     "Aturan ini ditegakkan di lapisan data, bukan di tiap halaman — supaya "
     "tidak ada halaman yang lupa memeriksanya.",
     ukuran=15, warna=REDUP)
catatan(s, "Kaitkan dengan risiko yang masih terbuka sekarang: tautan halaman "
           "redaksi sudah ada di footer semua halaman, tapi login belum "
           "dipasang. Itu ada di slide risiko.")

# =============================================================================
# 9 — Bukti
# =============================================================================
s = slide_baru()
judul_halaman(s, "163 pengecekan otomatis, semuanya lulus", "bukti")
teks(s, TEPI, Inches(1.9), ISI, Inches(0.4),
     "Dijalankan ulang 11 September 2026, bukan hasil lama",
     ukuran=14, warna=REDUP)
tabel(s, ["Perintah", "Hasil"], [
    ["npm run typecheck", ("PASS", HIJAU)],
    ["npm run build", ("PASS — 6 halaman", HIJAU)],
    ["npm run verify:repo", ("PASS — 37 pengecekan", HIJAU)],
    ["npm run verify:compare", ("PASS — 20 sama, 0 beda", HIJAU)],
    ["npm run verify:all", ("PASS — 106 pengecekan", HIJAU)],
    ["npm run lint", ("2 error, 5 warning", KUNING)],
    ["Uji database manual & pemeriksaan browser", ("NOT_RUN sejak PR #3", KUNING)],
], atas=Inches(2.4), lebar_kolom=[Inches(6.4), Inches(5.1)])
teks(s, TEPI, Inches(5.9), ISI, Inches(0.6),
     "Status yang boleh dipakai cuma PASS, FAIL, dan NOT_RUN. Yang belum "
     "dijalankan ditulis NOT_RUN, tidak dianggap lulus.",
     ukuran=15, warna=REDUP)
catatan(s, "Dua baris kuning sengaja ditampilkan. Kalau ditanya soal 2 error "
           "lint: keduanya ada di berkas yang bukan bagian lapisan data, dan "
           "sudah dikomunikasikan ke pemiliknya.")

# =============================================================================
# 10 — Kondisi sekarang
# =============================================================================
s = slide_baru()
judul_halaman(s, "Kondisi sekarang, apa adanya", "status")
tabel(s, ["Bagian", "Status", "Catatan"], [
    ["Lapisan data", ("Selesai", HIJAU), "5 tabel, 4 kumpulan fungsi, teruji"],
    ["Tampilan publik", ("Sebagian", KUNING), "Halaman ada, datanya masih dari dalam kode"],
    ["Halaman admin", ("Sebagian", KUNING), "Ada, tapi menyimpan ke browser, bukan database"],
    ["Login & kunci admin", ("Belum", MERAH), "Tabel akun siap, sistemnya belum dibuat"],
    ["Validasi & SEO", ("Belum", MERAH), "Belum ada sitemap, RSS, maupun validasi"],
    ["Deploy & pengecekan otomatis", ("Belum", MERAH), "Branch utama belum dikunci"],
    ["Dokumentasi", ("Selesai", HIJAU), "README, PRD, panduan, laporan, status"],
], atas=Inches(2.1), lebar_kolom=[Inches(3.6), Inches(1.9), Inches(6.0)])
catatan(s, "Jangan mempercantik slide ini. Justru kejujurannya yang membuat "
           "slide sebelumnya dipercaya.")

# =============================================================================
# 11 — Risiko
# =============================================================================
s = slide_baru()
judul_halaman(s, "Risiko dan keputusan yang menggantung", "perlu diputuskan")
poin(s, [
    ("Status komersial proyek belum jelas",
     "Vercel paket gratis melarang pemakaian komersial. Kalau ternyata untuk "
     "perusahaan, anggaran hosting harus dibahas sekarang"),
    ("Dua peran tim masih kosong",
     "Tanpa Orang 5, branch utama tidak terkunci dan tidak ada pengecekan otomatis"),
    ("Tautan halaman redaksi sudah publik di footer",
     "Login harus jalan SEBELUM halaman admin menyentuh database"),
    ("Halaman admin belum menyimpan ke database",
     "Artikel yang ditambah cuma terlihat di perangkat yang menambahkannya"),
], atas=Inches(2.0), jarak=Inches(1.15))
catatan(s, "Sampaikan risiko sebagai permintaan keputusan, bukan keluhan. "
           "Poin pertama yang paling mahal kalau baru ketahuan di minggu "
           "terakhir.")

# =============================================================================
# 12 — Langkah berikutnya
# =============================================================================
s = slide_baru()
judul_halaman(s, "Langkah berikutnya", "rencana")
tabel(s, ["Urutan", "Yang dikerjakan", "Siapa"], [
    ["1", "Kunci branch utama dan pasang pengecekan otomatis", "Orang 5"],
    ["2", "Bereskan 2 error lint supaya pengecekan bisa hijau", "Orang 2 & 3"],
    ["3", "Pasang login dan kunci halaman redaksi", "Orang 3"],
    ["4", "Sambungkan halaman admin ke database", "Orang 3 + Orang 1"],
    ["5", "Pindahkan halaman publik ke fungsi akses data", "Orang 2"],
    ["6", "Halaman kategori, pencarian, dan tombol muat lebih banyak", "Orang 2"],
    ["7", "Validasi, newsletter, sitemap, dan RSS", "Orang 4"],
], atas=Inches(2.1), lebar_kolom=[Inches(1.2), Inches(7.4), Inches(2.9)])
catatan(s, "Urutannya bukan asal. Nomor 3 harus sebelum nomor 4, karena "
           "tautan admin sudah publik.")

# =============================================================================
# 13 — Tim
# =============================================================================
s = slide_baru()
judul_halaman(s, "Pembagian tim", "siapa pegang apa")
tabel(s, ["Peran", "Wilayah", "Status"], [
    ["Orang 1 — Database", "prisma/, src/server/", ("Selesai", HIJAU)],
    ["Orang 2 — Tampilan Publik", "src/components/, halaman publik", ("Berjalan", KUNING)],
    ["Orang 3 — Admin & Login", "src/app/admin/", ("Berjalan", KUNING)],
    ["Orang 4 — Validasi & SEO", "validation/, services/", ("Belum ada orangnya", MERAH)],
    ["Orang 5 — Deploy & Testing", "workflows/, tests/", ("Belum ada orangnya", MERAH)],
], atas=Inches(2.1), lebar_kolom=[Inches(3.5), Inches(4.6), Inches(3.4)])
teks(s, TEPI, Inches(4.85), ISI, Inches(0.8),
     "Aturan yang menjaga lima orang tidak bertabrakan di repo sekecil ini: "
     "satu orang satu wilayah, dan mengubah wilayah orang lain harus lewat "
     "review pemiliknya.",
     ukuran=15, warna=REDUP)
catatan(s, "Kalau ditanya kenapa dua peran kosong: anggota organisasi ada 5, "
           "tapi dua orang belum menentukan bagian. Ini yang mau kami "
           "putuskan minggu ini.")

# =============================================================================
# 14 — Penutup
# =============================================================================
s = slide_baru(GELAP)
teks(s, TEPI, Inches(1.5), ISI, Inches(0.4), "PENUTUP", ukuran=12,
     warna=BIRU, tebal=True)
teks(s, TEPI, Inches(2.1), Inches(11.2), Inches(1.6),
     "Fondasinya sudah berdiri dan terbukti.\nYang tersisa: menyambungkannya.",
     ukuran=32, warna=PUTIH, tebal=True)
teks(s, TEPI, Inches(4.3), Inches(11.2), Inches(1.4),
     ["Dokumentasi lengkap ada di repo:",
      "STATUS.md · docs/PRD.md · docs/DATABASE.md · docs/LAPORAN-DATABASE.md",
      "github.com/article-website/news-times"],
     ukuran=15, warna=RGBColor(0xCB, 0xD5, 0xE1), spasi=Pt(6))
catatan(s, "Tutup dengan mengundang pertanyaan, dan arahkan ke STATUS.md untuk "
           "siapa pun yang ingin memeriksa sendiri.")


keluaran = Path(__file__).with_name("NewsTimes.pptx")
prs.save(keluaran)
print(f"Selesai: {keluaran}  ({len(prs.slides._sldIdLst)} slide)")
